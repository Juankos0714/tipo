from __future__ import annotations

import argparse
import logging
import numbers
import re
import unicodedata
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Protocol, Sequence

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

LOGGER = logging.getLogger(__name__)

ERROR = "ERROR"
SOSPECHOSO = "SOSPECHOSO"
OK = "OK"

RELLENOS = {
    ERROR: PatternFill("solid", fgColor="FFC7CE"),
    SOSPECHOSO: PatternFill("solid", fgColor="FFEB9C"),
}

HOJA_RESUMEN = "Resumen validacion"
HOJA_DETALLE = "Detalle validacion"
HOJA_TODAS = "Todas"


@dataclass(frozen=True)
class Columnas:
    factura_nota: str = "Fact./Nota"
    estado: str = "Estado"
    razon_social: str = "Razón social cliente factura"
    bruto: str = "Valor bruto"
    descuentos: str = "Valor descuentos"
    subtotal: str = "Valor subtotal"
    impuestos: str = "Valor impuestos"
    neto: str = "Valor neto"
    vendedor: str = "Nombre vendedor"
    credito: str = "Valor crédito"
    contado: str = "Valor contado"
    cajero: str = "Nombre cajero"
    efectivo: str = "Efectivo"
    cheque: str = "Cheque"
    tarjeta_credito: str = "Tarjeta crédito"
    tarjeta_debito: str = "Tarjeta débito"
    otros: str = "Otros"
    factura_electronica: str = "Factura electrónica"


@dataclass(frozen=True)
class Hallazgo:
    regla: str
    severidad: str
    detalle: str
    esperado: float | str
    encontrado: float | str


@dataclass(frozen=True)
class Registro:
    fila_excel: int
    factura: str
    razon_social: str
    estado: str
    factura_electronica: str
    cajero: str
    vendedor: str
    bruto: float
    descuentos: float
    subtotal: float
    impuestos: float
    neto: float
    credito: float
    contado: float
    efectivo: float
    cheque: float
    tarjeta_credito: float
    tarjeta_debito: float
    otros: float
    repeticiones: int

    @property
    def pagos(self) -> float:
        return (self.efectivo + self.cheque + self.tarjeta_credito
                + self.tarjeta_debito + self.otros)


@dataclass(frozen=True)
class Resultado:
    tabla: pd.DataFrame
    detalle: pd.DataFrame
    columnas: dict[str, str]


def normalizar(texto: str) -> str:
    descompuesto = unicodedata.normalize("NFD", str(texto))
    sin_tildes = "".join(c for c in descompuesto if unicodedata.category(c) != "Mn")
    sin_puntuacion = re.sub(r"[.,]", "", sin_tildes)
    return re.sub(r"\s+", " ", sin_puntuacion).strip().upper()


def a_texto(valor: object) -> str:
    return "" if pd.isna(valor) else str(valor).strip()


def a_numero(valor: object) -> float:
    if pd.isna(valor):
        return 0.0
    if isinstance(valor, numbers.Number):
        return float(valor)
    texto = re.sub(r"[$\s]", "", str(valor))
    if not texto:
        return 0.0
    if "," in texto:
        texto = texto.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(\.\d{3})+", texto):
        texto = texto.replace(".", "")
    try:
        return float(texto)
    except ValueError:
        return 0.0


def formato(numero: float) -> str:
    return f"{numero:,.2f}".rstrip("0").rstrip(".")


class Regla(Protocol):
    nombre: str
    severidad: str

    def evaluar(self, registro: Registro) -> list[Hallazgo]:
        ...


class ReglaRazonSocialVacia:
    nombre = "Razón social vacía"
    severidad = ERROR

    def evaluar(self, registro: Registro) -> list[Hallazgo]:
        if registro.razon_social:
            return []
        return [Hallazgo(self.nombre, self.severidad,
                         "La columna de razón social está vacía", "", "")]


class ReglaFacturaElectronica:
    nombre = "Factura electrónica no confirmada"
    severidad = ERROR

    def evaluar(self, registro: Registro) -> list[Hallazgo]:
        if normalizar(registro.factura_electronica) == "CONFIRMADOS":
            return []
        encontrado = registro.factura_electronica or "(vacío)"
        return [Hallazgo(self.nombre, self.severidad,
                         f"Valor en Factura electrónica: {encontrado}",
                         "Confirmados", encontrado)]


class ReglaEstado:
    nombre = "Estado distinto de Aprobadas"
    severidad = SOSPECHOSO

    def evaluar(self, registro: Registro) -> list[Hallazgo]:
        if normalizar(registro.estado) == "APROBADAS":
            return []
        return [Hallazgo(self.nombre, self.severidad,
                         f"Valor en Estado: {registro.estado or '(vacío)'}",
                         "Aprobadas", registro.estado or "(vacío)")]


class ReglaSubtotal:
    nombre = "Bruto - descuentos distinto de subtotal"
    severidad = ERROR

    def __init__(self, tolerancia: float) -> None:
        self._tolerancia = tolerancia

    def evaluar(self, registro: Registro) -> list[Hallazgo]:
        esperado = registro.bruto - registro.descuentos
        if abs(esperado - registro.subtotal) <= self._tolerancia:
            return []
        return [Hallazgo(self.nombre, self.severidad,
                         f"Bruto {formato(registro.bruto)} - descuentos {formato(registro.descuentos)}",
                         esperado, registro.subtotal)]


class ReglaNeto:
    nombre = "Subtotal + impuestos distinto de valor neto"
    severidad = ERROR

    def __init__(self, tolerancia: float) -> None:
        self._tolerancia = tolerancia

    def evaluar(self, registro: Registro) -> list[Hallazgo]:
        esperado = registro.subtotal + registro.impuestos
        if abs(esperado - registro.neto) <= self._tolerancia:
            return []
        return [Hallazgo(self.nombre, self.severidad,
                         f"Subtotal {formato(registro.subtotal)} + impuestos {formato(registro.impuestos)}",
                         esperado, registro.neto)]


class ReglaContadoCredito:
    nombre = "Contado + crédito distinto de valor neto"
    severidad = ERROR

    def __init__(self, tolerancia: float) -> None:
        self._tolerancia = tolerancia

    def evaluar(self, registro: Registro) -> list[Hallazgo]:
        esperado = registro.contado + registro.credito
        if abs(esperado - registro.neto) <= self._tolerancia:
            return []
        return [Hallazgo(self.nombre, self.severidad,
                         f"Contado {formato(registro.contado)} + crédito {formato(registro.credito)}",
                         esperado, registro.neto)]


class ReglaFormasDePago:
    nombre = "Formas de pago distintas del valor contado"
    severidad = ERROR

    def __init__(self, tolerancia: float) -> None:
        self._tolerancia = tolerancia

    def evaluar(self, registro: Registro) -> list[Hallazgo]:
        if abs(registro.pagos - registro.contado) <= self._tolerancia:
            return []
        detalle = (f"Efectivo {formato(registro.efectivo)} | Cheque {formato(registro.cheque)} | "
                   f"T.crédito {formato(registro.tarjeta_credito)} | "
                   f"T.débito {formato(registro.tarjeta_debito)} | Otros {formato(registro.otros)}")
        return [Hallazgo(self.nombre, self.severidad, detalle,
                         registro.contado, registro.pagos)]


class ReglaFacturaRepetida:
    nombre = "Fact./Nota repetida"
    severidad = ERROR

    def evaluar(self, registro: Registro) -> list[Hallazgo]:
        if not registro.factura or registro.repeticiones <= 1:
            return []
        return [Hallazgo(self.nombre, self.severidad,
                         f"Aparece {registro.repeticiones} veces", 1, registro.repeticiones)]


def reglas_por_defecto(tolerancia: float) -> list[Regla]:
    return [
        ReglaRazonSocialVacia(),
        ReglaFacturaElectronica(),
        ReglaEstado(),
        ReglaSubtotal(tolerancia),
        ReglaNeto(tolerancia),
        ReglaContadoCredito(tolerancia),
        ReglaFormasDePago(tolerancia),
        ReglaFacturaRepetida(),
    ]


class LectorExcel:
    def __init__(self, hoja: str | int = 0, fila_encabezado: int = 0) -> None:
        self._hoja = hoja
        self._fila_encabezado = fila_encabezado

    def leer(self, ruta: Path) -> pd.DataFrame:
        tabla = pd.read_excel(ruta, sheet_name=self._hoja, header=self._fila_encabezado)
        tabla.columns = [str(c).strip() for c in tabla.columns]
        return tabla.dropna(how="all").reset_index(drop=True)


class ValidadorVentas:
    def __init__(self, reglas: Sequence[Regla], columnas: Columnas,
                 fila_encabezado: int = 0) -> None:
        self._reglas = reglas
        self._columnas = columnas
        self._fila_encabezado = fila_encabezado

    def _resolver_columnas(self, tabla: pd.DataFrame) -> dict[str, str]:
        reales = {normalizar(c): c for c in tabla.columns}
        resueltas: dict[str, str] = {}
        faltantes: list[str] = []
        for campo, titulo in asdict(self._columnas).items():
            real = reales.get(normalizar(titulo))
            if real is None:
                faltantes.append(titulo)
            else:
                resueltas[campo] = real
        if faltantes:
            raise ValueError(f"Faltan columnas: {faltantes}. "
                             f"Columnas disponibles: {list(tabla.columns)}")
        return resueltas

    @staticmethod
    def _repeticiones(tabla: pd.DataFrame, columna: str) -> pd.Series:
        facturas = tabla[columna].map(a_texto)
        conteo = facturas[facturas != ""].value_counts()
        return facturas.map(conteo).fillna(0).astype(int)

    def _construir_registro(self, fila: pd.Series, c: dict[str, str],
                            indice: int, repeticiones: int) -> Registro:
        return Registro(
            fila_excel=indice + self._fila_encabezado + 2,
            factura=a_texto(fila[c["factura_nota"]]),
            razon_social=a_texto(fila[c["razon_social"]]),
            estado=a_texto(fila[c["estado"]]),
            factura_electronica=a_texto(fila[c["factura_electronica"]]),
            cajero=a_texto(fila[c["cajero"]]) or "(sin cajero)",
            vendedor=a_texto(fila[c["vendedor"]]) or "(sin vendedor)",
            bruto=a_numero(fila[c["bruto"]]),
            descuentos=a_numero(fila[c["descuentos"]]),
            subtotal=a_numero(fila[c["subtotal"]]),
            impuestos=a_numero(fila[c["impuestos"]]),
            neto=a_numero(fila[c["neto"]]),
            credito=a_numero(fila[c["credito"]]),
            contado=a_numero(fila[c["contado"]]),
            efectivo=a_numero(fila[c["efectivo"]]),
            cheque=a_numero(fila[c["cheque"]]),
            tarjeta_credito=a_numero(fila[c["tarjeta_credito"]]),
            tarjeta_debito=a_numero(fila[c["tarjeta_debito"]]),
            otros=a_numero(fila[c["otros"]]),
            repeticiones=repeticiones,
        )

    @staticmethod
    def _severidad(hallazgos: list[Hallazgo]) -> str:
        if not hallazgos:
            return OK
        if any(h.severidad == ERROR for h in hallazgos):
            return ERROR
        return SOSPECHOSO

    def validar(self, tabla: pd.DataFrame) -> Resultado:
        columnas = self._resolver_columnas(tabla)
        repeticiones = self._repeticiones(tabla, columnas["factura_nota"])
        severidades: list[str] = []
        observaciones: list[str] = []
        filas_detalle: list[dict[str, object]] = []

        for indice, fila in tabla.iterrows():
            registro = self._construir_registro(fila, columnas, int(indice),
                                                int(repeticiones.iloc[int(indice)]))
            hallazgos: list[Hallazgo] = []
            for regla in self._reglas:
                hallazgos.extend(regla.evaluar(registro))
            severidades.append(self._severidad(hallazgos))
            observaciones.append("; ".join(h.regla for h in hallazgos) or OK)
            for h in hallazgos:
                numerico = isinstance(h.esperado, float | int) and isinstance(h.encontrado, float | int)
                filas_detalle.append({
                    "Severidad": h.severidad,
                    "Fila Excel": registro.fila_excel,
                    "Fact./Nota": registro.factura,
                    "Razón social": registro.razon_social,
                    "Cajero": registro.cajero,
                    "Vendedor": registro.vendedor,
                    "Regla": h.regla,
                    "Detalle": h.detalle,
                    "Esperado": h.esperado,
                    "Encontrado": h.encontrado,
                    "Diferencia": h.encontrado - h.esperado if numerico else "",
                })

        resultado = tabla.copy()
        resultado["Severidad"] = severidades
        resultado["Observaciones"] = observaciones
        columnas_detalle = ["Severidad", "Fila Excel", "Fact./Nota", "Razón social", "Cajero",
                            "Vendedor", "Regla", "Detalle", "Esperado", "Encontrado", "Diferencia"]
        detalle = pd.DataFrame(filas_detalle, columns=columnas_detalle)
        return Resultado(resultado, detalle, columnas)


class ConstructorResumen:
    def __init__(self, reglas: Sequence[Regla]) -> None:
        self._reglas = reglas

    @staticmethod
    def _porcentaje(parte: int, total: int) -> float:
        return round(parte * 100 / total, 1) if total else 0.0

    def construir(self, resultado: Resultado) -> list[tuple[str, pd.DataFrame]]:
        tabla = resultado.tabla
        total = len(tabla)
        errores = int((tabla["Severidad"] == ERROR).sum())
        sospechosos = int((tabla["Severidad"] == SOSPECHOSO).sum())

        general = pd.DataFrame({
            "Indicador": ["Filas revisadas", "Filas con ERROR", "Filas SOSPECHOSAS", "Filas OK"],
            "Cantidad": [total, errores, sospechosos, total - errores - sospechosos],
            "% de filas": [100.0, self._porcentaje(errores, total),
                           self._porcentaje(sospechosos, total),
                           self._porcentaje(total - errores - sospechosos, total)],
        })

        conteo = resultado.detalle["Regla"].value_counts()
        por_regla = pd.DataFrame({
            "Regla": [r.nombre for r in self._reglas],
            "Severidad": [r.severidad for r in self._reglas],
            "Cantidad": [int(conteo.get(r.nombre, 0)) for r in self._reglas],
        })
        por_regla["% de filas"] = [self._porcentaje(c, total) for c in por_regla["Cantidad"]]

        return [
            ("RESUMEN DE VALIDACIÓN", general),
            ("Hallazgos por regla", por_regla),
            ("Por cajero", self._por_persona(tabla, resultado.columnas["cajero"], "Cajero")),
            ("Por vendedor", self._por_persona(tabla, resultado.columnas["vendedor"], "Vendedor")),
        ]

    def _por_persona(self, tabla: pd.DataFrame, columna: str, etiqueta: str) -> pd.DataFrame:
        nombres = tabla[columna].map(a_texto).replace("", f"(sin {etiqueta.lower()})")
        agrupado = (
            pd.DataFrame({etiqueta: nombres, "con_hallazgos": tabla["Severidad"] != OK})
            .groupby(etiqueta)["con_hallazgos"]
            .agg(Facturas="count", **{"Con hallazgos": "sum"})
            .reset_index()
        )
        agrupado["% con hallazgos"] = [
            self._porcentaje(int(a), int(b))
            for a, b in zip(agrupado["Con hallazgos"], agrupado["Facturas"])
        ]
        return agrupado.sort_values("Con hallazgos", ascending=False).reset_index(drop=True)


class EscritorInforme:
    def escribir(self, resultado: Resultado, bloques: list[tuple[str, pd.DataFrame]],
                 ruta: Path) -> None:
        with pd.ExcelWriter(ruta, engine="openpyxl") as escritor:
            self._escribir_resumen(escritor, bloques)
            resultado.detalle.to_excel(escritor, sheet_name=HOJA_DETALLE, index=False)
            resultado.tabla.to_excel(escritor, sheet_name=HOJA_TODAS, index=False)
            self._formatear_tabla(escritor.sheets[HOJA_DETALLE])
            self._formatear_tabla(escritor.sheets[HOJA_TODAS])
            self._colorear_severidad(escritor.sheets[HOJA_DETALLE], "Severidad", 1)
            self._colorear_severidad(escritor.sheets[HOJA_TODAS], "Severidad", 2)
            self._ajustar_anchos(escritor.sheets[HOJA_RESUMEN])

    @staticmethod
    def _escribir_resumen(escritor: pd.ExcelWriter, bloques: list[tuple[str, pd.DataFrame]]) -> None:
        fila = 0
        for titulo, tabla in bloques:
            pd.DataFrame([[titulo]]).to_excel(escritor, sheet_name=HOJA_RESUMEN,
                                              startrow=fila, index=False, header=False)
            tabla.to_excel(escritor, sheet_name=HOJA_RESUMEN, startrow=fila + 1, index=False)
            hoja = escritor.sheets[HOJA_RESUMEN]
            hoja.cell(row=fila + 1, column=1).font = Font(bold=True, size=12)
            for columna in range(1, len(tabla.columns) + 1):
                hoja.cell(row=fila + 2, column=columna).font = Font(bold=True)
            fila += len(tabla) + 3

    def _formatear_tabla(self, hoja) -> None:
        for celda in hoja[1]:
            celda.font = Font(bold=True)
        hoja.freeze_panes = "A2"
        hoja.auto_filter.ref = hoja.dimensions
        self._ajustar_anchos(hoja)

    @staticmethod
    def _ajustar_anchos(hoja, maximo: int = 60) -> None:
        for columna in hoja.columns:
            largo = max((len(str(c.value)) for c in columna if c.value is not None), default=8)
            hoja.column_dimensions[get_column_letter(columna[0].column)].width = min(largo + 2, maximo)

    @staticmethod
    def _colorear_severidad(hoja, titulo: str, columnas_a_pintar: int) -> None:
        indices = {c.value: c.column for c in hoja[1]}
        if titulo not in indices:
            return
        inicio = indices[titulo]
        for fila in range(2, hoja.max_row + 1):
            relleno = RELLENOS.get(hoja.cell(row=fila, column=inicio).value)
            if relleno is None:
                continue
            for desplazamiento in range(columnas_a_pintar):
                hoja.cell(row=fila, column=inicio + desplazamiento).fill = relleno


def imprimir_resumen(nombre: str, resultado: Resultado, reglas: Sequence[Regla],
                     ejemplos: int) -> None:
    tabla = resultado.tabla
    errores = int((tabla["Severidad"] == ERROR).sum())
    sospechosos = int((tabla["Severidad"] == SOSPECHOSO).sum())
    LOGGER.info("%s | filas=%d errores=%d sospechosos=%d hallazgos=%d",
                nombre, len(tabla), errores, sospechosos, len(resultado.detalle))
    for regla in reglas:
        encontrados = resultado.detalle[resultado.detalle["Regla"] == regla.nombre]
        LOGGER.info("  %s [%s]: %d", regla.nombre, regla.severidad, len(encontrados))
        for _, h in encontrados.head(ejemplos).iterrows():
            LOGGER.info("     fila %s | %s | esperado %s | encontrado %s | %s",
                        h["Fila Excel"], h["Fact./Nota"], h["Esperado"], h["Encontrado"], h["Detalle"])


def recopilar_archivos(rutas: Sequence[Path]) -> list[Path]:
    archivos: list[Path] = []
    for ruta in rutas:
        if ruta.is_dir():
            for patron in ("*.xlsx", "*.xls"):
                archivos.extend(p for p in sorted(ruta.glob(patron)) if not p.name.startswith("~$"))
        elif ruta.is_file():
            archivos.append(ruta)
        else:
            LOGGER.warning("No existe: %s", ruta)
    return archivos


def construir_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Valida reportes de ventas y facturación en Excel.")
    parser.add_argument("rutas", nargs="+", type=Path, help="Archivos .xlsx/.xls o carpetas")
    parser.add_argument("--hoja", default=0, help="Nombre o posición de la hoja (0 = primera)")
    parser.add_argument("--fila-encabezado", type=int, default=0, help="Fila de títulos, desde 0")
    parser.add_argument("--tolerancia", type=float, default=1.0, help="Diferencia permitida en pesos")
    parser.add_argument("--ejemplos", type=int, default=5, help="Ejemplos por regla en pantalla")
    parser.add_argument("--salida", type=Path, default=Path("resultados"))
    return parser


def main() -> int:
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
    args = construir_parser().parse_args()
    hoja = int(args.hoja) if str(args.hoja).isdigit() else args.hoja
    args.salida.mkdir(parents=True, exist_ok=True)

    reglas = reglas_por_defecto(args.tolerancia)
    lector = LectorExcel(hoja, args.fila_encabezado)
    validador = ValidadorVentas(reglas, Columnas(), args.fila_encabezado)
    constructor = ConstructorResumen(reglas)
    escritor = EscritorInforme()

    archivos = recopilar_archivos(args.rutas)
    if not archivos:
        LOGGER.error("No se encontraron archivos para procesar")
        return 1

    fallidos = 0
    for archivo in archivos:
        try:
            resultado = validador.validar(lector.leer(archivo))
            destino = args.salida / f"{archivo.stem}_validado.xlsx"
            escritor.escribir(resultado, constructor.construir(resultado), destino)
            imprimir_resumen(archivo.name, resultado, reglas, args.ejemplos)
            LOGGER.info("Informe: %s", destino)
        except Exception:
            LOGGER.exception("Falló el archivo %s", archivo)
            fallidos += 1
    return 1 if fallidos else 0


if __name__ == "__main__":
    raise SystemExit(main())
